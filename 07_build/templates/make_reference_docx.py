#!/usr/bin/env python3
"""Erzeugt die Word-Formatvorlage (reference.docx) für den Pandoc-Export.

Die Vorlage legt das verbindliche Layout des Lehrbuchs fest: Seitenformat,
Schriften, Überschriften, Kästen, Tabellen- und Beschriftungsstile sowie
Kopf- und Fußzeile. Sie wird bei jedem Build neu erzeugt, damit Kopfzeilen-
texte (Dokumenttitel, Stand) je Dokument gesetzt werden können.

Schriften: Überschriften/Tabellen/Kästen in Calibri, Fließtext in Cambria.
Beide sind in jeder Office-Installation vorhanden; im Build-Container werden
die metrisch identischen freien Schriften Carlito und Caladea verwendet.

Aufruf:
    python3 make_reference_docx.py --out reference.docx \
        --kopf-links "Kurztitel des Buchs" --kopf-rechts "Stand TT.MM.JJJJ"
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

warnings.filterwarnings("ignore", category=UserWarning, module="docx")

# ---------------------------------------------------------------------------
# Gestaltungskonstanten (zentral, damit das Layout an einer Stelle änderbar ist)
# ---------------------------------------------------------------------------
FONT_HEAD = "Calibri"      # Überschriften, Tabellen, Kästen, Beschriftungen
FONT_BODY = "Cambria"      # Fließtext
FONT_MONO = "Consolas"     # Code

NAVY = "1B365D"            # Primärfarbe (Überschriften, Tabellenkopf)
ACCENT = "2A7AB0"          # Akzent (H3, Linien, Links)
TEXT = "1A1A1A"            # Fließtext
GRAY = "5F6B7A"            # Nebentexte (Kopf-/Fußzeile, Metadaten)
RULE = "C9D3E0"            # feine Linien

# Kastentypen: (Linienfarbe, Füllfarbe) – wird auch vom Post-Processing genutzt
KASTEN_FARBEN = {
    "merke": ("2A7AB0", "EAF2FA"),
    "fachleute": ("C55A11", "FDF0E6"),
    "praxis": ("2E7D6B", "E8F4F0"),
    "achtung": ("B7791F", "FFF6E0"),
    "beispiel": ("1B365D", "F1F3F6"),
    "lernziele": ("5B4B8A", "F1EEF8"),
    "definition": ("5F6B7A", "F5F6F8"),
    "stand": ("8A1C3C", "FBEFF2"),
}

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
    "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
    "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange",
]


def _local(tag: str) -> str:
    return tag.split("}")[1] if "}" in tag else tag


def insert_ordered(parent, child, order):
    """Fügt child schemakonform (Reihenfolge laut OOXML) in parent ein."""
    tag = _local(child.tag)
    idx = order.index(tag)
    for existing in list(parent):
        etag = _local(existing.tag)
        if etag == tag:
            parent.replace(existing, child)
            return child
    for existing in list(parent):
        etag = _local(existing.tag)
        if etag in order and order.index(etag) > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def pandoc_bin() -> str:
    exe = shutil.which("pandoc")
    if exe:
        return exe
    try:
        import pypandoc  # type: ignore

        return pypandoc.get_pandoc_path()
    except Exception as exc:  # pragma: no cover
        sys.exit(f"pandoc nicht gefunden: {exc}")


def get_style(doc, name: str, kind=WD_STYLE_TYPE.PARAGRAPH, base: str | None = None):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, kind)
        if base:
            style.base_style = doc.styles[base]
    style.hidden = False
    style.quick_style = True
    return style


def set_font(style, name=None, size=None, bold=None, italic=None, color=None, small_caps=None):
    font = style.font
    if name:
        font.name = name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(att)) is not None:
                del rfonts.attrib[qn(att)]
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(att), name)
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color:
        font.color.rgb = RGBColor.from_string(color)
    if small_caps is not None:
        font.small_caps = small_caps


def set_para(style, before=None, after=None, line=None, align=None, keep_next=None,
             keep_lines=None, left=None, right=None, first=None, page_break=None):
    pf = style.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if align is not None:
        pf.alignment = align
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if keep_lines is not None:
        pf.keep_together = keep_lines
    if left is not None:
        pf.left_indent = Cm(left)
    if right is not None:
        pf.right_indent = Cm(right)
    if first is not None:
        pf.first_line_indent = Cm(first)
    if page_break is not None:
        pf.page_break_before = page_break
    pf.widow_control = True


def set_border(style, side: str, color: str, size_eighths: int, space_pt: int = 4, val="single"):
    ppr = style.element.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        insert_ordered(ppr, pbdr, PPR_ORDER)
    order = ["top", "left", "bottom", "right", "between", "bar"]
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(size_eighths))
    el.set(qn("w:space"), str(space_pt))
    el.set(qn("w:color"), color)
    insert_ordered(pbdr, el, order)


def set_outline_level(style, level: int):
    ppr = style.element.get_or_add_pPr()
    el = OxmlElement("w:outlineLvl")
    el.set(qn("w:val"), str(level))
    insert_ordered(ppr, el, PPR_ORDER)


def add_tab(style, pos_cm: float, align="right", leader=None):
    ppr = style.element.get_or_add_pPr()
    tabs = ppr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        insert_ordered(ppr, tabs, PPR_ORDER)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), align)
    if leader:
        tab.set(qn("w:leader"), leader)
    tab.set(qn("w:pos"), str(int(pos_cm * 567)))
    tabs.append(tab)


def add_field(paragraph, instr: str, size=8.5, color=GRAY, bold=False):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FONT_HEAD)
    rpr.append(rfonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rpr.append(col)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def style_run(run, size=8.5, color=GRAY, bold=False, italic=False, small_caps=False):
    run.font.name = FONT_HEAD
    rfonts = run._r.get_or_add_rPr().find(qn("w:rFonts"))
    for att in ("w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FONT_HEAD)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    run.font.italic = italic
    run.font.small_caps = small_caps


def set_doc_defaults(doc):
    styles_el = doc.styles.element
    rpr = styles_el.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr"))
    if rpr is None:
        return
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        for att in list(rfonts.attrib):
            del rfonts.attrib[att]
        for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(att), FONT_BODY)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "de-DE")
    lang.set(qn("w:eastAsia"), "de-DE")
    lang.set(qn("w:bidi"), "ar-SA")


SETTINGS_ORDER = [
    "writeProtection", "view", "zoom", "removePersonalInformation", "removeDateAndTime",
    "doNotDisplayPageBoundaries", "displayBackgroundShape", "printPostScriptOverText",
    "printFractionalCharacterWidth", "printFormsData", "embedTrueTypeFonts", "embedSystemFonts",
    "saveSubsetFonts", "saveFormsData", "mirrorMargins", "alignBordersAndEdges",
    "bordersDoNotSurroundHeader", "bordersDoNotSurroundFooter", "gutterAtTop",
    "hideSpellingErrors", "hideGrammaticalErrors", "activeWritingStyle", "proofState",
    "formsDesign", "attachedTemplate", "linkStyles", "stylePaneFormatFilter",
    "stylePaneSortMethod", "documentType", "mailMerge", "revisionView", "trackRevisions",
    "doNotTrackMoves", "doNotTrackFormatting", "documentProtection", "autoFormatOverride",
    "styleLockTheme", "styleLockQFSet", "defaultTabStop", "autoHyphenation",
    "consecutiveHyphenLimit", "hyphenationZone", "doNotHyphenateCaps", "showEnvelope",
    "summaryLength", "clickAndTypeStyle", "defaultTableStyle", "evenAndOddHeaders",
    "bookFoldRevPrinting", "bookFoldPrinting", "bookFoldPrintingSheets",
    "drawingGridHorizontalSpacing", "drawingGridVerticalSpacing",
    "displayHorizontalDrawingGridEvery", "displayVerticalDrawingGridEvery",
    "doNotUseMarginsForDrawingGridOrigin", "drawingGridHorizontalOrigin",
    "drawingGridVerticalOrigin", "doNotShadeFormData", "noPunctuationKerning",
    "characterSpacingControl", "printTwoOnOne", "strictFirstAndLastChars", "noLineBreaksAfter",
    "noLineBreaksBefore", "savePreviewPicture", "doNotValidateAgainstSchema", "saveInvalidXml",
    "ignoreMixedContent", "alwaysShowPlaceholderText", "doNotDemarcateInvalidXml",
    "saveXmlDataOnly", "useXSLTWhenSaving", "saveThroughXslt", "showXMLTags",
    "alwaysMergeEmptyNamespace", "updateFields", "hdrShapeDefaults", "footnotePr", "endnotePr",
    "compat", "docVars", "rsids", "mathPr", "attachedSchema", "themeFontLang",
    "clrSchemeMapping", "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures",
    "forceUpgrade", "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
    "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator",
]


def configure_settings(doc):
    """Automatische Silbentrennung (deutsch) aktivieren."""
    settings = doc.settings.element

    def upsert(tag, val=None):
        el = OxmlElement(f"w:{tag}")
        if val is not None:
            el.set(qn("w:val"), val)
        insert_ordered(settings, el, SETTINGS_ORDER)

    upsert("autoHyphenation", "true")
    upsert("consecutiveHyphenLimit", "3")
    upsert("hyphenationZone", "425")
    upsert("doNotHyphenateCaps", "true")


def build(out: Path, kopf_links: str, kopf_rechts: str, fuss_links: str):
    tmp = Path(tempfile.mkdtemp()) / "pandoc-reference.docx"
    subprocess.run([pandoc_bin(), "-o", str(tmp), "--print-default-data-file", "reference.docx"],
                   check=True)
    doc = Document(str(tmp))
    set_doc_defaults(doc)

    # --- Grundschrift -----------------------------------------------------
    normal = get_style(doc, "Normal")
    set_font(normal, FONT_BODY, 10.5, color=TEXT)
    set_para(normal, before=0, after=6, line=1.2)

    for name in ("Body Text", "First Paragraph"):
        st = get_style(doc, name)
        set_font(st, FONT_BODY, 10.5, color=TEXT)
        set_para(st, before=0, after=6, line=1.2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    compact = get_style(doc, "Compact")
    set_font(compact, FONT_BODY, 10.5, color=TEXT)
    set_para(compact, before=0, after=2, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)

    # --- Titelblock (Deckblatt) ------------------------------------------------
    title = get_style(doc, "Title")
    set_font(title, FONT_HEAD, 28, bold=True, color=NAVY)
    set_para(title, before=96, after=10, line=1.0, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
    set_border(title, "top", NAVY, 36, space_pt=18)
    ppr = title.element.get_or_add_pPr()
    for el in ppr.findall(qn("w:pBdr")):
        for side in ("bottom", "left", "right"):
            for b in el.findall(qn(f"w:{side}")):
                el.remove(b)

    subtitle = get_style(doc, "Subtitle")
    set_font(subtitle, FONT_HEAD, 15, bold=False, italic=False, color=ACCENT)
    set_para(subtitle, before=0, after=28, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)

    author = get_style(doc, "Author")
    set_font(author, FONT_HEAD, 11.5, bold=True, color=TEXT)
    set_para(author, before=0, after=2, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)

    date = get_style(doc, "Date")
    set_font(date, FONT_HEAD, 11, color=GRAY)
    set_para(date, before=0, after=18, line=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)

    for name, size in (("Abstract", 10), ("Abstract Title", 10)):
        try:
            st = doc.styles[name]
            set_font(st, FONT_HEAD, size, color=GRAY)
        except KeyError:
            pass

    meta = get_style(doc, "Deckblatt Meta", base="Normal")
    set_font(meta, FONT_HEAD, 10, color=TEXT)
    set_para(meta, before=0, after=3, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_tab(meta, 4.2, align="left")

    hinweis = get_style(doc, "Deckblatt Hinweis", base="Normal")
    set_font(hinweis, FONT_HEAD, 8.5, italic=True, color=GRAY)
    set_para(hinweis, before=24, after=0, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_border(hinweis, "top", RULE, 6, space_pt=6)

    # --- Überschriften -----------------------------------------------------
    specs = {
        1: dict(size=18, color=NAVY, before=6, after=14, italic=False),
        2: dict(size=14, color=NAVY, before=18, after=6, italic=False),
        3: dict(size=12, color=ACCENT, before=12, after=4, italic=False),
        4: dict(size=11, color=NAVY, before=10, after=3, italic=True),
        5: dict(size=10.5, color=GRAY, before=8, after=2, italic=True),
    }
    for level, sp in specs.items():
        st = get_style(doc, f"Heading {level}")
        set_font(st, FONT_HEAD, sp["size"], bold=True, italic=sp["italic"], color=sp["color"])
        set_para(st, before=sp["before"], after=sp["after"], line=1.05,
                 align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, keep_lines=True)
    set_border(get_style(doc, "Heading 1"), "bottom", ACCENT, 12, space_pt=6)

    toc_head = get_style(doc, "TOC Heading")
    set_font(toc_head, FONT_HEAD, 18, bold=True, color=NAVY)
    set_para(toc_head, before=6, after=18, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
    set_outline_level(toc_head, 9)
    set_border(toc_head, "bottom", ACCENT, 12, space_pt=6)

    # Verzeichnisstile (Word-Name "toc 1" etc.; LibreOffice: "Contents 1")
    text_width_cm = 21.0 - 2.5 - 2.3
    for level, (size, bold, before, indent) in {
        1: (10.5, True, 8, 0.0), 2: (10, False, 1, 0.6), 3: (9.5, False, 0, 1.4)
    }.items():
        style_id = f"TOC{level}"
        existing = doc.styles.element.find(
            f"{qn('w:style')}[@{qn('w:styleId')}='{style_id}']")
        if existing is not None:
            doc.styles.element.remove(existing)
        el = OxmlElement("w:style")
        el.set(qn("w:type"), "paragraph")
        el.set(qn("w:styleId"), style_id)
        name = OxmlElement("w:name")
        name.set(qn("w:val"), f"toc {level}")
        el.append(name)
        based = OxmlElement("w:basedOn")
        based.set(qn("w:val"), "Normal")
        el.append(based)
        nxt = OxmlElement("w:next")
        nxt.set(qn("w:val"), "Normal")
        el.append(nxt)
        prio = OxmlElement("w:uiPriority")
        prio.set(qn("w:val"), "39")
        el.append(prio)
        el.append(OxmlElement("w:unhideWhenUsed"))
        doc.styles.element.append(el)
        try:
            st = doc.styles[f"toc {level}"]
        except KeyError:
            continue
        set_font(st, FONT_HEAD, size, bold=bold, color=NAVY if level == 1 else TEXT)
        set_para(st, before=before, after=2, line=1.1, left=indent)
        add_tab(st, text_width_cm, align="right", leader="dot")

    # --- Beschriftungen, Abbildungen --------------------------------------
    for name in ("Table Caption", "Image Caption", "Caption"):
        try:
            st = doc.styles[name]
        except KeyError:
            st = get_style(doc, name, base="Normal")
        set_font(st, FONT_HEAD, 9.5, italic=False, color=NAVY)
        set_para(st, before=8 if name == "Table Caption" else 4,
                 after=4 if name == "Table Caption" else 12, line=1.1,
                 align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=(name == "Table Caption"))
    for name in ("Figure", "Captioned Figure"):
        try:
            st = doc.styles[name]
            set_para(st, before=6, after=2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        except KeyError:
            pass

    # --- Tabellenabsätze ----------------------------------------------------
    t_text = get_style(doc, "Tabelle Text", base="Normal")
    set_font(t_text, FONT_HEAD, 9, color=TEXT)
    set_para(t_text, before=1, after=1, line=1.05, align=WD_ALIGN_PARAGRAPH.LEFT)
    t_head = get_style(doc, "Tabelle Kopf", base="Tabelle Text")
    set_font(t_head, FONT_HEAD, 9, bold=True, color="FFFFFF")
    set_para(t_head, before=1, after=1, line=1.05, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)

    # --- Kästen -------------------------------------------------------------
    k_text = get_style(doc, "Kasten Text", base="Normal")
    set_font(k_text, FONT_HEAD, 10, color=TEXT)
    set_para(k_text, before=0, after=4, line=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    k_title = get_style(doc, "Kasten Titel", base="Kasten Text")
    set_font(k_title, FONT_HEAD, 10, bold=True, color=NAVY, small_caps=True)
    set_para(k_title, before=0, after=3, line=1.1, keep_next=True)
    marker = get_style(doc, "Kasten Marker", base="Normal")
    set_font(marker, FONT_HEAD, 2)

    # --- Zitate, Fußnoten, Code, Links -------------------------------------
    block = get_style(doc, "Block Text")
    set_font(block, FONT_BODY, 10, italic=False, color="3A4452")
    set_para(block, before=4, after=8, line=1.15, left=0.6, right=0.6,
             align=WD_ALIGN_PARAGRAPH.LEFT)
    set_border(block, "left", RULE, 18, space_pt=8)

    for name in ("Footnote Text",):
        try:
            st = doc.styles[name]
            set_font(st, FONT_BODY, 8.5, color=TEXT)
            set_para(st, before=0, after=2, line=1.05, align=WD_ALIGN_PARAGRAPH.LEFT)
        except KeyError:
            pass

    try:
        code = doc.styles["Source Code"]
        set_font(code, FONT_MONO, 8.5, color=TEXT)
    except KeyError:
        pass
    try:
        verb = doc.styles["Verbatim Char"]
        set_font(verb, FONT_MONO, 9, color=TEXT)
    except KeyError:
        pass
    try:
        link = doc.styles["Hyperlink"]
        link.font.color.rgb = RGBColor.from_string(ACCENT)
        link.font.underline = False
    except KeyError:
        pass

    spacer = get_style(doc, "Abstand", base="Normal")
    set_font(spacer, FONT_BODY, 4)
    set_para(spacer, before=0, after=0, line=1.0)

    # --- Seite, Kopf- und Fußzeile --------------------------------------------
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.6)
    sec.bottom_margin = Cm(2.3)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.3)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.0)
    sec.different_first_page_header_footer = True

    width = Cm(text_width_cm)
    hdr = sec.header
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]
    p.text = ""
    p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_after = Pt(0)
    style_run(p.add_run(kopf_links), size=8.5, color=GRAY, small_caps=True)
    style_run(p.add_run("\t" + kopf_rechts), size=8.5, color=GRAY)
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "4"), ("w:color", RULE)):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    insert_ordered(ppr, pbdr, PPR_ORDER)

    ftr = sec.footer
    ftr.is_linked_to_previous = False
    p = ftr.paragraphs[0]
    p.text = ""
    p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    style_run(p.add_run(fuss_links), size=8, color=GRAY)
    style_run(p.add_run("\tSeite "), size=8.5, color=GRAY)
    add_field(p, "PAGE", size=8.5, color=NAVY, bold=True)
    style_run(p.add_run(" von "), size=8.5, color=GRAY)
    add_field(p, "NUMPAGES", size=8.5, color=GRAY)

    # erste Seite (Deckblatt) ohne Kopf-/Fußzeile
    sec.first_page_header.is_linked_to_previous = False
    sec.first_page_footer.is_linked_to_previous = False

    configure_settings(doc)
    doc.core_properties.language = "de-DE"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=Path(__file__).with_name("reference.docx"))
    ap.add_argument("--kopf-links", default="Lehrbuch")
    ap.add_argument("--kopf-rechts", default="Lehrbuchprojekt")
    ap.add_argument("--fuss-links", default="Entwurf – keine Rechts-, Steuer- oder Anlageberatung")
    args = ap.parse_args()
    path = build(args.out, args.kopf_links, args.kopf_rechts, args.fuss_links)
    print(f"reference.docx geschrieben: {path}")


if __name__ == "__main__":
    main()
